"""
Testes de ponta a ponta com os 14 e-mails do desafio.
Se algum destes falhar, o robô deixou de bater com o gabarito.
"""

from openpyxl import load_workbook

from conftest import DATA_DESAFIO
from robo import planilha
from robo.processador import (DUPLICADA, EXCECAO, IGNORADO, JA_LANCADA,
                              processar)


def _rodar(emails, arquivo, pasta_anexos=None):
    return processar(emails, arquivo, DATA_DESAFIO, pasta_anexos,
                     log=lambda _: None)


def _soma(contas):
    return round(sum(c["Valor (R$)"] for c in contas), 2)


def test_nivel_1_so_xml(emails, planilha_tmp):
    res = _rodar(emails, planilha_tmp)
    nfe = [c for c in res.contas if c["Tipo"] == "NF-e"]
    assert len(nfe) == 6
    assert _soma(nfe) == 6340.57


def test_nivel_2_xml_e_pdf(emails, planilha_tmp):
    res = _rodar(emails, planilha_tmp)
    anexos = [c for c in res.contas if c["Tipo"] in ("NF-e", "Fatura PDF")]
    assert len(anexos) == 8
    assert _soma(anexos) == 8282.75
    assert res.contar(DUPLICADA) == 1
    assert res.contar(EXCECAO) == 1


def test_nivel_3_com_corpo_do_email(emails, planilha_tmp):
    res = _rodar(emails, planilha_tmp)
    assert len(res.abertas) == 9
    assert res.total == 8631.75
    assert res.contar(IGNORADO) == 3


def test_nenhum_email_some_sem_explicacao(emails, planilha_tmp):
    res = _rodar(emails, planilha_tmp)
    arquivos_no_log = {item["Arquivo"] for item in res.log}
    assert arquivos_no_log == {e.nome for e in emails}


def test_idempotencia(emails, planilha_tmp):
    _rodar(emails, planilha_tmp)
    segunda = _rodar(emails, planilha_tmp)
    assert segunda.novas == []
    assert segunda.contar(JA_LANCADA) == 9
    assert segunda.contar(DUPLICADA) == 1   # o reenvio segue duplicado
    assert len(segunda.contas) == 9


def test_conta_paga_sai_do_total(emails, planilha_tmp):
    _rodar(emails, planilha_tmp)

    wb = load_workbook(planilha_tmp)
    ws = wb[planilha.ABA_CONTAS]
    linhas = list(ws.iter_rows())
    titulos = next(r for r in linhas if r[0].value == "Status")
    cab = [c.value for c in titulos]
    for linha in linhas[linhas.index(titulos) + 1:]:
        if linha[cab.index("Documento")].value == "NF-e 1074":
            linha[cab.index("Pago em")].value = DATA_DESAFIO
    wb.save(planilha_tmp)

    res = _rodar(emails, planilha_tmp)
    paga = [c for c in res.contas if c["Documento"] == "NF-e 1074"]
    assert paga[0]["Status"] == planilha.PAGA
    assert res.total == round(8631.75 - 1850.00, 2)


def test_anexos_organizados(emails, planilha_tmp, tmp_path):
    anexos = tmp_path / "anexos"
    res = _rodar(emails, planilha_tmp, anexos)
    assert res.anexos_salvos == 9
    assert (anexos / "Nuvem Hosting Ltda" / "2026-09" / "NFe_1000.xml"
            ).exists()
    # Rodar de novo não duplica arquivos
    assert _rodar(emails, planilha_tmp, anexos).anexos_salvos == 0
