"""
Robô de Contas a Pagar - Nível 1

Lê todos os e-mails .eml da pasta de entrada, encontra os que trazem NF-e em
XML anexado, extrai os dados da nota e gera a planilha contas_a_pagar.xlsx.
E-mails que não são cobrança são ignorados (e registrados no log do terminal).

Uso:
    python robo_nivel1.py
"""

import email
import email.policy
import os
import re
import xml.etree.ElementTree as ET
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# ---------------------------------------------------------------------------
# Configuração
# ---------------------------------------------------------------------------
PASTA_BASE = Path(__file__).parent


def carregar_env(caminho: Path) -> None:
    """Lê linhas CHAVE=valor do .env para os.environ (sem sobrescrever)."""
    if not caminho.exists():
        return
    for linha in caminho.read_text(encoding="utf-8").splitlines():
        linha = linha.strip()
        if not linha or linha.startswith("#") or "=" not in linha:
            continue
        chave, valor = linha.split("=", 1)
        os.environ.setdefault(chave.strip(), valor.strip().strip('"\''))


carregar_env(PASTA_BASE / ".env")

PASTA_EMAILS = PASTA_BASE / os.getenv("PASTA_EMAILS",
                                      "dados/caixa_de_entrada")
ARQUIVO_SAIDA = PASTA_BASE / os.getenv("ARQUIVO_SAIDA", "contas_a_pagar.xlsx")
# Data fixa, para o resultado ser reprodutível
DATA_REFERENCIA = date.fromisoformat(os.getenv("DATA_REFERENCIA",
                                               "2026-09-20"))

# Namespace oficial da NF-e. Sem ele, o ElementTree não encontra as tags.
NS = {"nfe": "http://www.portalfiscal.inf.br/nfe"}

COLUNAS = [
    "Fornecedor",
    "CNPJ",
    "Número NF",
    "Data de Emissão",
    "Valor Total (R$)",
    "Vencimento",
    "Arquivo de Origem",
]


# ---------------------------------------------------------------------------
# 1. Ler os e-mails
# ---------------------------------------------------------------------------
def ler_email(caminho: Path) -> email.message.EmailMessage:
    """Abre um arquivo .eml e devolve o objeto de e-mail já interpretado."""
    with open(caminho, "rb") as f:
        return email.message_from_binary_file(f, policy=email.policy.default)


def anexos_xml(msg: email.message.EmailMessage) -> list[tuple[str, bytes]]:
    """Devolve a lista de anexos XML do e-mail como (nome_do_arquivo, conteúdo)."""
    encontrados = []
    for parte in msg.iter_attachments():
        nome = parte.get_filename() or ""
        tipo = parte.get_content_type()
        if nome.lower().endswith(".xml") or tipo in ("application/xml", "text/xml"):
            encontrados.append((nome, parte.get_payload(decode=True)))
    return encontrados


# ---------------------------------------------------------------------------
# 2. Extrair os dados da NF-e
# ---------------------------------------------------------------------------
def texto(raiz: ET.Element, caminho: str) -> str | None:
    """Busca uma tag (com namespace) e devolve o texto, ou None se não existir."""
    elemento = raiz.find(caminho, NS)
    return elemento.text.strip() if elemento is not None and elemento.text else None


def formatar_cnpj(cnpj: str) -> str:
    """14 dígitos -> 12.345.678/0001-90"""
    return f"{cnpj[:2]}.{cnpj[2:5]}.{cnpj[5:8]}/{cnpj[8:12]}-{cnpj[12:]}"


def extrair_nfe(conteudo: bytes) -> dict:
    """
    Lê o XML de uma NF-e e devolve um dicionário com os campos da conta.
    Lança ValueError se o XML estiver quebrado ou não for uma NF-e.
    """
    # Alguns emissores mandam "&" cru (ex.: "Silva & Filhos"), o que é XML
    # inválido. Escapamos todo "&" que não inicia uma entidade (&amp; &#38; ...).
    conteudo = re.sub(rb"&(?!(?:[a-zA-Z]+|#\d+|#x[0-9a-fA-F]+);)", b"&amp;", conteudo)
    try:
        raiz = ET.fromstring(conteudo)
    except ET.ParseError as erro:
        raise ValueError(f"XML inválido/corrompido ({erro})")

    inf = raiz.find(".//nfe:infNFe", NS)
    if inf is None:
        raise ValueError("XML não é uma NF-e (tag infNFe ausente)")

    cnpj = texto(inf, "nfe:emit/nfe:CNPJ")
    fornecedor = texto(inf, "nfe:emit/nfe:xNome")
    numero = texto(inf, "nfe:ide/nfe:nNF")
    emissao = texto(inf, "nfe:ide/nfe:dhEmi")      # 2026-09-02T09:00:00-03:00
    valor = texto(inf, "nfe:total/nfe:ICMSTot/nfe:vNF")
    vencimento = texto(inf, "nfe:cobr/nfe:dup/nfe:dVenc")  # 1ª duplicata

    obrigatorios = {"CNPJ": cnpj, "fornecedor": fornecedor, "número": numero,
                    "emissão": emissao, "valor": valor}
    faltando = [campo for campo, v in obrigatorios.items() if not v]
    if faltando:
        raise ValueError(f"campos ausentes no XML: {', '.join(faltando)}")

    return {
        "Fornecedor": fornecedor,
        "CNPJ": formatar_cnpj(cnpj),
        "Número NF": int(numero),
        "Data de Emissão": datetime.fromisoformat(emissao).date(),
        "Valor Total (R$)": float(valor),
        "Vencimento": date.fromisoformat(vencimento) if vencimento else None,
    }


# ---------------------------------------------------------------------------
# 3. Gerar a planilha
# ---------------------------------------------------------------------------
def gerar_planilha(contas: list[dict], destino: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Contas a Pagar"

    ws.append(COLUNAS)
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="305496")

    for conta in contas:
        ws.append([conta[c] for c in COLUNAS])

    # Linha de total
    ultima = ws.max_row
    ws.append(["TOTAL", None, None, None, f"=SUM(E2:E{ultima})"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=5).font = Font(bold=True)

    # Formatos de exibição
    for linha in ws.iter_rows(min_row=2):
        linha[3].number_format = "DD/MM/YYYY"
        linha[4].number_format = '"R$" #,##0.00'
        linha[5].number_format = "DD/MM/YYYY"

    larguras = [34, 20, 11, 16, 17, 13, 50]
    for i, largura in enumerate(larguras):
        ws.column_dimensions[chr(ord("A") + i)].width = largura

    wb.save(destino)


# ---------------------------------------------------------------------------
# Orquestração
# ---------------------------------------------------------------------------
def main() -> None:
    arquivos = sorted(PASTA_EMAILS.glob("*.eml"))
    print(f"Data de referência: {DATA_REFERENCIA:%d/%m/%Y}")
    print(f"{len(arquivos)} e-mails encontrados em {PASTA_EMAILS}\n")

    contas: list[dict] = []
    ja_lancadas: set[tuple[str, int]] = set()   # chave: (CNPJ, número da nota)

    for arquivo in arquivos:
        try:
            msg = ler_email(arquivo)
        except Exception as erro:
            print(f"[ERRO]      {arquivo.name}: não foi possível ler o e-mail ({erro})")
            continue

        xmls = anexos_xml(msg)
        if not xmls:
            print(f"[IGNORADO]  {arquivo.name}: sem NF-e em XML")
            continue

        for nome_anexo, conteudo in xmls:
            try:
                conta = extrair_nfe(conteudo)
            except ValueError as erro:
                print(f"[ERRO]      {arquivo.name} / {nome_anexo}: {erro}")
                continue

            chave = (conta["CNPJ"], conta["Número NF"])
            if chave in ja_lancadas:
                print(f"[DUPLICADA] {arquivo.name}: NF {conta['Número NF']} de "
                      f"{conta['Fornecedor']} já lançada")
                continue

            ja_lancadas.add(chave)
            conta["Arquivo de Origem"] = arquivo.name
            contas.append(conta)
            print(f"[OK]        {arquivo.name}: NF {conta['Número NF']} - "
                  f"{conta['Fornecedor']} - R$ {conta['Valor Total (R$)']:.2f}")

    contas.sort(key=lambda c: (c["Vencimento"] or date.max, c["Fornecedor"]))
    try:
        gerar_planilha(contas, ARQUIVO_SAIDA)
    except PermissionError:
        print(f"\n[ERRO] Não foi possível salvar {ARQUIVO_SAIDA.name}: "
              "feche a planilha no Excel e rode de novo.")
        return

    total = sum(c["Valor Total (R$)"] for c in contas)
    total_br = f"{total:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    print(f"\n{len(contas)} contas lançadas | Total a pagar: R$ {total_br}")
    print(f"Planilha gerada: {ARQUIVO_SAIDA}")


if __name__ == "__main__":
    main()
